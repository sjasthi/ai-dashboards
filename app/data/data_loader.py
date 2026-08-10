import os
import re

import pandas as pd

from app.data.sheet_scan import scan_rows

EXCEL_EXTENSIONS = {".xlsx", ".xls"}

_UNNAMED_RE = re.compile(r"^Unnamed: \d+$")


class DataLoader:

    def __init__(self):
        self.files = []
        # Table name -> the uploaded file it came from. "Orders (sales).xlsx" can't
        # be parsed back reliably, since sheet names and workbook stems may both
        # contain parentheses, so the mapping is recorded while it is still known.
        self.origins = {}
        # Uploaded filename -> {"kind", "sheet_count"}: what the file turned out to
        # be once opened. Keyed per upload rather than per table, so a workbook
        # whose sheets were all deselected or empty still reports how many it holds.
        self.sources = {}
        # Table display name -> human-readable notes about a structural repair
        # made while loading it (e.g. a shifted header row, a relabeled unnamed
        # column). Folded into that table's quality_flags by SummaryGenerator so
        # the LLM sees them - see summary_builder.profile_all_files.
        self.load_warnings = {}
        # Uploaded filename -> error string, for a file that raised while loading.
        # Recorded rather than left to blow up add_files, so one bad file (a
        # locked Office companion file, a corrupt upload) doesn't sink a batch
        # that otherwise loaded fine.
        self.load_failures = {}

    def add_files(self, file_paths, selections=None):
        """Add files, dispatching by extension. CSVs are read with encoding
        detection; Excel workbooks contribute one entry per worksheet.

        selections optionally narrows which worksheets are loaded, as
        {uploaded filename: [sheet names]}. A workbook absent from the mapping
        contributes every sheet, so callers that don't care can omit it entirely.
        Filtering here means excluded sheets never reach profiling, the prompt or
        the tables map.

        One file failing to load - a locked Office companion file, a corrupt
        upload - is recorded in load_failures rather than raised, so it can't
        take down a batch where every other file is fine.
        """
        for file_path in file_paths:
            try:
                ext = os.path.splitext(file_path)[1].lower()
                if ext in EXCEL_EXTENSIONS:
                    self._add_excel(file_path, selections)
                else:
                    df = self._read_csv_with_encoding(file_path)
                    self._note(file_path, "csv")
                    self._record(file_path, df, os.path.basename(file_path))
            except Exception as e:
                name = os.path.basename(file_path)
                self.load_failures[name] = f"{type(e).__name__}: {e}"
                print(f"[DataLoader] Skipping {name}: {type(e).__name__}: {e}")

    def _record(self, name, df, source):
        self.files.append((name, df))
        self.origins[os.path.basename(name)] = source

    def _note(self, file_path, kind, sheet_count=None):
        """What one uploaded file turned out to be. sheet_count stays None for a
        CSV: it has no worksheets at all, which is a different fact from a
        one-sheet workbook."""
        self.sources[os.path.basename(file_path)] = {
            "kind": kind, "sheet_count": sheet_count,
        }

    def _add_excel(self, file_path, selections=None):
        try:
            sheets = pd.read_excel(file_path, sheet_name=None)
        except ValueError:
            # Handle CSV mislabeled as xlsx.
            df = self._read_csv_with_encoding(file_path)
            self._note(file_path, "csv")
            self._record(file_path, df, os.path.basename(file_path))
            return
        base = os.path.basename(file_path)
        stem, ext = os.path.splitext(base)

        # notes[sheet_name] collects human-readable descriptions of any repair
        # made to that sheet, surfaced later via self.load_warnings.
        notes = {}

        # A sheet whose real header isn't row 0 (e.g. a blank spacer row above
        # it) reads as all "Unnamed: N" columns above - detect that from the
        # same raw-cell scan workbook_probe.py uses for the upload preview, so
        # the two never disagree about where a sheet's data starts, and re-read
        # just the affected sheets with the correct header.
        header_rows = _detect_header_rows(file_path, ext)
        reread = {
            name: offset for name, offset in header_rows.items()
            if offset and name in sheets
        }
        if reread:
            # Closed explicitly (not just left to the ExcelFile destructor): on
            # Windows an unclosed handle keeps the temp file's zip archive open,
            # which makes the caller's later shutil.rmtree(temp_dir) fail with
            # PermissionError: [WinError 32].
            with pd.ExcelFile(file_path) as xls:
                for sheet_name, offset in reread.items():
                    sheets[sheet_name] = pd.read_excel(xls, sheet_name=sheet_name, header=offset)
                    notes.setdefault(sheet_name, []).append(
                        f"header row detected {offset} row(s) down in the source "
                        f"file (blank row(s) above it were skipped)"
                    )

        # Before any filtering: this is the workbook's own size, not how much of
        # it was asked for. read_excel(sheet_name=None) has already read them all.
        self._note(file_path, "excel", len(sheets))

        chosen = (selections or {}).get(base)

        # Collect survivors before naming anything: both the selection filter and
        # the empty-sheet skip can drop the count below two, and the parenthesised
        # name is only warranted when a sheet really does have siblings.
        kept = []
        for sheet_name, df in sheets.items():
            if chosen is not None and sheet_name not in chosen:
                continue
            if df.empty:
                continue
            # Strip only string headers; Excel headers can be numeric/blank
            df.columns = df.columns.map(lambda c: c.strip() if isinstance(c, str) else c)
            unnamed_notes = _label_unnamed_columns(df, sheet_name)
            if unnamed_notes:
                notes.setdefault(sheet_name, []).extend(unnamed_notes)
            kept.append((sheet_name, df))

        multi = len(kept) > 1
        for sheet_name, df in kept:
            # Lead with the sheet name: for a multi-sheet workbook the sheet is the
            # meaningful table identity, and relationship detection matches foreign
            # keys against the *start* of the file stem (see summary_builder). The
            # workbook stem in parens keeps names unique across workbooks.
            display = f"{sheet_name} ({stem}){ext}" if multi else base
            self._record(display, df, base)
            if notes.get(sheet_name):
                self.load_warnings[display] = notes[sheet_name]
            print(f"Loaded {file_path} (sheet: {sheet_name})")

    def _read_csv_with_encoding(self, file_path):
        """Try to read CSV with multiple encoding options."""
        encodings = ['utf-8', 'utf-16', 'latin-1', 'iso-8859-1', 'cp1252']

        for encoding in encodings:
            try:
                df = pd.read_csv(file_path, encoding=encoding)
                df.columns = df.columns.str.strip()
                print(f"Loaded {file_path} (encoding: {encoding})")
                return df
            except (UnicodeDecodeError, UnicodeError):
                continue

        # If all fail, raise informative error
        raise ValueError(
            f"Could not read {file_path} with any supported encoding. "
            f"Tried: {encodings}"
        )

    def tables(self):
        """Map table name -> DataFrame, keyed the same way profiling names files
        for the LLM. Worksheets aren't files on disk, so this is the only way to
        get a sheet's data back after loading."""
        return {os.path.basename(name): df for name, df in self.files}


def _detect_header_rows(file_path, ext):
    """Sheet name -> 0-based header row offset for pandas' read_excel(header=...).

    Uses the same raw-cell scan workbook_probe.py uses for the upload preview
    (sheet_scan.scan_rows), so DataLoader and the preview can never disagree
    about where a sheet's data actually starts. 0 for a fully blank sheet -
    pandas still needs some offset, and an empty result is skipped downstream
    by the df.empty check regardless of which row it "used" as header.
    """
    if ext == ".xls":
        import xlrd

        book = xlrd.open_workbook(file_path, on_demand=True)
        try:
            offsets = {}
            for name in book.sheet_names():
                sheet = book.sheet_by_name(name)
                result = scan_rows(sheet.row_values(r) for r in range(sheet.nrows))
                offsets[name] = max(result.header_row - 1, 0)
                book.unload_sheet(name)
            return offsets
        finally:
            book.release_resources()

    import openpyxl

    wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
    try:
        return {
            ws.title: max(scan_rows(ws.iter_rows(values_only=True)).header_row - 1, 0)
            for ws in wb.worksheets
        }
    finally:
        wb.close()


def _label_unnamed_columns(df, sheet_name):
    """Rename pandas' "Unnamed: N" placeholders (a header cell with no text) to
    an unambiguous synthetic name, in place, and return one note per renamed
    column.

    Deliberately doesn't try to guess a real name from the column's data - that
    would be fabrication. A clearly-synthetic label plus a note that reaches the
    LLM (via quality_flags) is the honest version of "this column has no name
    in the source file."
    """
    notes = []
    existing = {str(c) for c in df.columns}
    new_columns = []
    for position, col in enumerate(df.columns, start=1):
        if not (isinstance(col, str) and _UNNAMED_RE.match(col)):
            new_columns.append(col)
            continue
        candidate = f"Column_{position}"
        suffix = 1
        while candidate in existing:
            suffix += 1
            candidate = f"Column_{position}_{suffix}"
        existing.add(candidate)
        new_columns.append(candidate)
        notes.append(
            f"{sheet_name}: column {position} has no header name in the source "
            f"file (labeled '{candidate}')"
        )
    if notes:
        df.columns = new_columns
    return notes
